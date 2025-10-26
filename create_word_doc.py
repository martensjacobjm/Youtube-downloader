#!/usr/bin/env python3
"""
Skript för att skapa ett Word-dokument från XML-innehåll
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def create_word_from_xml(xml_file_path, output_docx_path):
    """
    Läser XML från Word document.xml och skapar ett nytt .docx-dokument

    OBS: Detta är en förenklad version som fungerar med ren text.
    För fullständig XML-parsing behövs mer komplex kod.
    """

    # Skapa nytt Word-dokument
    doc = Document()

    print(f"Läser XML från: {xml_file_path}")

    # Läs XML-innehållet
    with open(xml_file_path, 'r', encoding='utf-8') as f:
        xml_content = f.read()

    # Extrahera text från <w:t>-taggar (Word text-element)
    text_pattern = r'<w:t[^>]*>(.*?)</w:t>'
    texts = re.findall(text_pattern, xml_content)

    print(f"Hittade {len(texts)} textstycken")

    # Lägg till texten i dokumentet
    for text in texts:
        if text.strip():  # Hoppa över tomma strängar
            # Avkoda XML-entiteter
            text = text.replace('&amp;', '&')
            text = text.replace('&lt;', '<')
            text = text.replace('&gt;', '>')
            text = text.replace('&quot;', '"')

            # Lägg till som paragraf
            p = doc.add_paragraph(text)

    # Spara dokumentet
    doc.save(output_docx_path)
    print(f"✅ Word-dokument sparat: {output_docx_path}")

def create_formatted_testur_doc(output_path):
    """
    Skapar ett formaterat TesTur-dokument baserat på innehållet
    """
    doc = Document()

    # Titel
    title = doc.add_heading('KAPITEL 3.2: TESTUR ENERGY 10-TUMS TURBIN', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Författare
    author = doc.add_paragraph('Jacob Martens')
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.runs[0].bold = True

    # Version och datum
    version = doc.add_paragraph('Version 5.0')
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date = doc.add_paragraph('2025-10-26')
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Sidbrytning
    doc.add_page_break()

    # Innehållsförteckning
    doc.add_heading('Innehållsförteckning', 1)
    doc.add_paragraph('[Automatisk innehållsförteckning skapas i Word]')

    doc.add_page_break()

    # Huvudinnehåll
    doc.add_heading('KAPITEL 3.2: TESTUR ENERGY 10-TUMS TURBIN', 1)

    doc.add_paragraph(
        'TesTur Energy utvecklade under tidigt 2020-tal en kommersiell 10-tums (254 mm) '
        'Tesla-turbin för lågtryckstilämpningar. Turbinen har genomgått omfattande testning '
        'och dokumentation, med både dynamometer-tester vid högt tryck och praktiska '
        'effekttester vid lägre tryck.'
    )

    # Lägg till fler sektioner här...
    doc.add_heading('3.2.1 TURBINMODELL', 2)
    doc.add_paragraph(
        "TesTur Energy's 10-tums turbin utvecklades som första kommersiella implementering "
        "av Tesla-turbinprincipen för moderna tillämpningar."
    )

    # Spara
    doc.save(output_path)
    print(f"✅ Formaterat dokument skapat: {output_path}")

if __name__ == "__main__":
    import sys

    print("=" * 60)
    print("Word-dokument skapare från XML")
    print("=" * 60)

    # Alternativ 1: Om du har document.xml
    if len(sys.argv) > 1 and sys.argv[1] == "from-xml":
        xml_path = "word/document.xml"
        output_path = "TesTur_Kapitel_3.2.docx"
        create_word_from_xml(xml_path, output_path)

    # Alternativ 2: Skapa formaterat dokument från scratch
    else:
        output_path = "TesTur_Kapitel_3.2_Formatted.docx"
        create_formatted_testur_doc(output_path)

